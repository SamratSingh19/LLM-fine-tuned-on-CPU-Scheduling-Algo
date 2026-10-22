import os, glob, pickle
import numpy as np
import faiss
from sentence_transformers import SentenceTransformer

def extract_pdf(path):
    import fitz
    doc = fitz.open(path)
    return "\n".join(page.get_text() for page in doc)

def extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)

def extract_docx(path):
    import docx
    doc = docx.Document(path)
    return "\n".join(p.text for p in doc.paragraphs)

def extract_txt(path):
    return open(path, encoding="utf-8", errors="ignore").read()

EXTRACTORS = {".pdf": extract_pdf, ".pptx": extract_pptx,
              ".docx": extract_docx, ".txt": extract_txt}

def load_and_chunk(folder="data/raw_docs", chunk_size=300, overlap=50):
    chunks = []
    for path in glob.glob(f"{folder}/*"):
        ext = os.path.splitext(path)[1].lower()
        if ext not in EXTRACTORS:
            continue
        text = EXTRACTORS[ext](path)
        words = text.split()
        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append({"text": chunk, "source": os.path.basename(path)})
    return chunks

def build_index(chunks, model_name="BAAI/bge-small-en-v1.5"):
    model = SentenceTransformer(model_name)
    embeddings = model.encode([c["text"] for c in chunks], normalize_embeddings=True,
                               show_progress_bar=True)
    index = faiss.IndexFlatIP(embeddings.shape[1])
    index.add(np.array(embeddings).astype("float32"))

    os.makedirs("data/processed", exist_ok=True)
    faiss.write_index(index, "data/processed/faiss.index")
    pickle.dump(chunks, open("data/processed/chunks.pkl", "wb"))
    print(f"Indexed {len(chunks)} chunks")

def search(query, k=3, model_name="BAAI/bge-small-en-v1.5"):
    model = SentenceTransformer(model_name)
    index = faiss.read_index("data/processed/faiss.index")
    chunks = pickle.load(open("data/processed/chunks.pkl", "rb"))
    q_emb = model.encode([query], normalize_embeddings=True).astype("float32")
    _, idxs = index.search(q_emb, k)
    return [chunks[i] for i in idxs[0]]

if __name__ == "__main__":
    chunks = load_and_chunk()
    build_index(chunks)
    print(search("What is Round Robin scheduling?"))